import sqlite3
import csv
import io
import secrets
from datetime import datetime, timezone

from fastapi import HTTPException
from pydantic import ValidationError

from app.repositories import product_repository
from app.models.schemas import ProductInput


class ProductService:
    @staticmethod
    def list_products(user, search: str, category: str, low_stock: bool):
        products = product_repository.list_products(user["tenant_id"], search, category, low_stock)
        if user["role"] != "admin" and not user.get("permissions", {}).get("view_purchase_costs", False):
            for product in products:
                product.pop("cost", None)
        return products

    @staticmethod
    def list_categories(user):
        return product_repository.list_categories(user["tenant_id"])

    @staticmethod
    def _clean(data):
        values = data.model_dump()
        values["name"] = values["name"].strip()
        values["brand"] = values["brand"].strip()
        values["category"] = values["category"].strip() or "General"
        values["sku"] = values["sku"].strip()
        if not values["sku"]:
            values["sku"] = f"CP{secrets.token_hex(5).upper()}"
        values["size"] = values["size"].strip()
        values["color"] = values["color"].strip()
        values["unit"] = values["unit"].strip() or "each"
        values["expiry_date"] = values["expiry_date"].strip()
        return values

    @classmethod
    def create_product(cls, user, data):
        try:
            return product_repository.create_product(
                user["tenant_id"], cls._clean(data), datetime.now(timezone.utc).isoformat()
            )
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=409, detail="That barcode / SKU is already in use.")

    @classmethod
    def import_products(cls, user, filename: str, contents: bytes):
        try:
            if filename.lower().endswith(".csv"):
                rows = list(csv.DictReader(io.StringIO(contents.decode("utf-8-sig"))))
            elif filename.lower().endswith(".xlsx"):
                from openpyxl import load_workbook

                workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
                sheet = workbook.active
                values = sheet.iter_rows(values_only=True)
                headers = [str(value or "").strip().lower() for value in next(values, ())]
                rows = [
                    dict(zip(headers, row))
                    for row in values
                    if any(value is not None for value in row)
                ]
            else:
                raise HTTPException(status_code=415, detail="Upload a .csv or .xlsx file.")
        except (UnicodeDecodeError, csv.Error, ValueError):
            raise HTTPException(status_code=422, detail="Could not read this file. Use UTF-8 CSV or XLSX.")
        if not rows:
            raise HTTPException(status_code=422, detail="The uploaded file has no product rows.")
        if len(rows) > 5000:
            raise HTTPException(status_code=422, detail="Import up to 5,000 products per file.")

        valid_products = []
        errors = []
        seen_skus = set()
        for row_number, row in enumerate(rows, start=2):
            normalized = {str(key).strip().lower(): value for key, value in row.items() if key is not None}
            try:
                product = ProductInput.model_validate(normalized)
                cleaned = cls._clean(product)
                if cleaned["sku"] in seen_skus:
                    errors.append({"row": row_number, "error": f"Duplicate SKU {cleaned['sku']} in this file."})
                else:
                    valid_products.append(cleaned)
                    seen_skus.add(cleaned["sku"])
            except (ValidationError, AttributeError, TypeError, ValueError) as exc:
                errors.append({"row": row_number, "error": str(exc)})
        if errors:
            raise HTTPException(status_code=422, detail={"message": "No rows imported; fix the listed rows and try again.", "errors": errors})
        try:
            imported = product_repository.import_products(
                user["tenant_id"], valid_products, datetime.now(timezone.utc).isoformat()
            )
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=409, detail="An imported SKU already exists in this store. No rows were imported.")
        return {"imported": imported}

    @classmethod
    def update_product(cls, user, product_id: int, data):
        values = cls._clean(data)
        if user["role"] != "admin" and not user.get("permissions", {}).get("view_purchase_costs", False):
            current = product_repository.get_product(user["tenant_id"], product_id)
            if current:
                values["cost"] = current["cost"]
        try:
            product = product_repository.update_product(
                user["tenant_id"], product_id, values, user["id"], datetime.now(timezone.utc).isoformat()
            )
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=409, detail="That barcode / SKU is already in use.")
        if not product:
            raise HTTPException(status_code=404, detail="Product not found.")
        return product

    @staticmethod
    def delete_product(user, product_id: int):
        if not product_repository.delete_product(user["tenant_id"], product_id):
            raise HTTPException(status_code=404, detail="Product not found.")
        return {"deleted": True}

    @staticmethod
    def adjust_stock(user, data):
        result = product_repository.adjust_stock(
            user["tenant_id"], user["id"], data.product_id, data.quantity, data.reason.strip(),
            datetime.now(timezone.utc).isoformat(),
        )
        if result is False:
            raise HTTPException(status_code=404, detail="Product not found.")
        if result is None:
            raise HTTPException(status_code=409, detail="Stock adjustment would make inventory negative.")
        return {"adjusted": True}
